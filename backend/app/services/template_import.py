"""Registry-driven reader for flat interval spreadsheets.

A ``tabular_intervals`` template (stored in ``import_profiles.mapping``) describes a
regular table: which sheet, which row holds the headers, which header feeds which
platform field, and how the borehole code is found. The reader turns such a file
into the same dataset shape ``normalize_excel_workbook`` produces, so the existing
merge path consumes it unchanged.

Built-in layouts (CTSJ, PBH) keep their hand-written readers in ``excel_import``.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from app.services.excel_import import clean, normalize_lithology, to_float

TEMPLATE_KIND = "tabular_intervals"
SCHEMA_VERSION = 2

# Platform fields a template may map. "lithology" plus a depth pair are required.
FIELD_KEYS = (
    "lithology",
    "logged_color",
    "structural_features",
    "recovery",
    "recovery_percent",
    "rqd",
    "seam_name",
    "remark",
    "grain_size",
    "core_dip",
)
BOREHOLE_SOURCES = ("column", "selected")
NORMALIZE_RULES = ("none", "upper", "zero_pad_number:2", "zero_pad_number:3")
RQD_UNITS = ("percent", "fraction")
SAMPLE_ROW_LIMIT = 25
MAX_HEADER_ROW = 50


def is_tabular_template(mapping: dict | None) -> bool:
    return isinstance(mapping, dict) and mapping.get("kind") == TEMPLATE_KIND


def header_key(value) -> str:
    """Case- and whitespace-insensitive header identity."""
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


# --------------------------------------------------------------------------- validation


def validate_template(mapping: dict) -> dict:
    """Return a normalized copy of ``mapping`` or raise ValueError listing every problem."""
    if not isinstance(mapping, dict):
        raise ValueError("Template mapping must be a JSON object.")
    errors: list[str] = []

    if mapping.get("kind") != TEMPLATE_KIND:
        errors.append(f'"kind" must be "{TEMPLATE_KIND}".')

    detect = mapping.get("detect") or {}
    if not isinstance(detect, dict):
        errors.append('"detect" must be an object.')
        detect = {}
    sheet = detect.get("sheet", "first")
    if not isinstance(sheet, str) or not sheet.strip():
        errors.append('"detect.sheet" must be "first" or a sheet name.')

    header_row = mapping.get("header_row")
    if not isinstance(header_row, int) or not 1 <= header_row <= MAX_HEADER_ROW:
        errors.append(f'"header_row" must be a whole number from 1 to {MAX_HEADER_ROW}.')
        header_row = 1
    data_start_row = mapping.get("data_start_row", header_row + 1)
    if not isinstance(data_start_row, int) or data_start_row <= header_row:
        errors.append('"data_start_row" must be a whole number after "header_row".')
        data_start_row = header_row + 1

    depth = mapping.get("depth") or {}
    if not isinstance(depth, dict) or not _is_ref(depth.get("from")):
        errors.append('"depth.from" must name the From-depth column.')
        depth = depth if isinstance(depth, dict) else {}
    has_to = _is_ref(depth.get("to"))
    has_thickness = _is_ref(depth.get("thickness"))
    if has_to == has_thickness:
        errors.append('Map exactly one of "depth.to" or "depth.thickness".')

    fields = mapping.get("fields") or {}
    if not isinstance(fields, dict):
        errors.append('"fields" must be an object.')
        fields = {}
    unknown = sorted(set(fields) - set(FIELD_KEYS))
    if unknown:
        errors.append(f"Unknown field(s): {', '.join(unknown)}. Allowed: {', '.join(FIELD_KEYS)}.")
    if not _is_ref(fields.get("lithology")):
        errors.append('"fields.lithology" must name the lithology column.')
    for key, ref in fields.items():
        if key == "remark" and isinstance(ref, list):
            if not ref or not all(_is_ref(item) for item in ref):
                errors.append('"fields.remark" list must contain column names.')
        elif ref not in (None, "") and not _is_ref(ref):
            errors.append(f'"fields.{key}" must be a column name.')

    borehole = mapping.get("borehole") or {"source": "selected"}
    if not isinstance(borehole, dict) or borehole.get("source") not in BOREHOLE_SOURCES:
        errors.append(f'"borehole.source" must be one of: {", ".join(BOREHOLE_SOURCES)}.')
        borehole = {"source": "selected"}
    if borehole.get("source") == "column" and not _is_ref(borehole.get("column")):
        errors.append('"borehole.column" must name the borehole-code column.')
    if borehole.get("normalize", "none") not in NORMALIZE_RULES:
        errors.append(f'"borehole.normalize" must be one of: {", ".join(NORMALIZE_RULES)}.')

    units = mapping.get("units") or {}
    if units.get("rqd", "percent") not in RQD_UNITS:
        errors.append(f'"units.rqd" must be one of: {", ".join(RQD_UNITS)}.')

    end = mapping.get("end") or {}
    stop_after = end.get("stop_after_blank_rows", 3) if isinstance(end, dict) else 3
    if not isinstance(stop_after, int) or stop_after < 1:
        errors.append('"end.stop_after_blank_rows" must be a whole number of at least 1.')
        stop_after = 3

    if errors:
        raise ValueError(" ".join(errors))

    required = detect.get("required_headers")
    if not required:
        required = _default_required_headers(depth, fields, borehole)
    if not isinstance(required, list) or not all(isinstance(item, str) and item.strip() for item in required):
        raise ValueError('"detect.required_headers" must be a list of column names.')

    return {
        "schema_version": SCHEMA_VERSION,
        "kind": TEMPLATE_KIND,
        "detect": {"sheet": sheet.strip(), "required_headers": [item.strip() for item in required]},
        "header_row": header_row,
        "data_start_row": data_start_row,
        "borehole": {
            "source": borehole["source"],
            **({"column": _ref_text(borehole["column"])} if borehole["source"] == "column" else {}),
            "normalize": borehole.get("normalize", "none"),
        },
        "depth": {
            "from": _ref_text(depth["from"]),
            **({"to": _ref_text(depth["to"])} if has_to else {"thickness": _ref_text(depth["thickness"])}),
        },
        "fields": {
            key: ([_ref_text(item) for item in ref] if isinstance(ref, list) else _ref_text(ref))
            for key, ref in fields.items()
            if ref not in (None, "", [])
        },
        "units": {"depth": "m", "rqd": units.get("rqd", "percent")},
        "end": {"stop_after_blank_rows": stop_after},
        "status": mapping.get("status", "active"),
        "version": mapping.get("version", 1),
    }


def _is_ref(value) -> bool:
    return isinstance(value, str) and bool(value.strip())


def _ref_text(value: str) -> str:
    return value.strip()


def _default_required_headers(depth: dict, fields: dict, borehole: dict) -> list[str]:
    refs = [depth.get("from"), depth.get("to") or depth.get("thickness"), fields.get("lithology")]
    if borehole.get("source") == "column":
        refs.insert(0, borehole.get("column"))
    return [ref.strip() for ref in refs if _is_ref(ref)]


# --------------------------------------------------------------------------- borehole codes


def normalize_borehole_code(value, rule: str = "none") -> str | None:
    text = clean(value)
    if text is None:
        return None
    code = re.sub(r"\s+", "", str(text)).upper()
    if rule.startswith("zero_pad_number:"):
        width = int(rule.split(":", 1)[1])
        match = re.match(r"^(.*?)(\d+)$", code)
        if match:
            code = f"{match.group(1)}{match.group(2).zfill(width)}"
    return code


def borehole_codes_match(file_code, borehole_code: str, rule: str) -> bool:
    left = normalize_borehole_code(file_code, rule)
    right = normalize_borehole_code(borehole_code, rule)
    return left is not None and left == right


# --------------------------------------------------------------------------- workbook access


def inspect_workbook(path: Path, max_rows: int = SAMPLE_ROW_LIMIT) -> dict:
    """First rows of every sheet, for the template builder's header picker."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(min_row=1, max_row=max_rows, values_only=True):
                rows.append([_cell_json(value) for value in row])
            width = max((len(row) for row in rows), default=0)
            sheets.append(
                {
                    "name": worksheet.title,
                    "max_row": worksheet.max_row,
                    "rows": [row + [None] * (width - len(row)) for row in rows],
                }
            )
        return {"sheets": sheets}
    finally:
        workbook.close()


def _cell_json(value):
    if value is None or isinstance(value, (int, float, str, bool)):
        return value
    return str(value)


def _select_sheet(workbook, sheet: str):
    if sheet == "first":
        return workbook.worksheets[0]
    for worksheet in workbook.worksheets:
        if header_key(worksheet.title) == header_key(sheet):
            return worksheet
    raise ValueError(f'Sheet "{sheet}" was not found. Sheets: {", ".join(workbook.sheetnames)}.')


def _read_header_index(worksheet, header_row: int) -> tuple[dict[str, int], list[str]]:
    headers: list[str] = []
    for row in worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        headers = ["" if value is None else str(value).strip() for value in row]
    index: dict[str, int] = {}
    for position, header in enumerate(headers):
        key = header_key(header)
        if key and key not in index:
            index[key] = position
    return index, headers


def template_matches(path: Path, template: dict) -> bool:
    """True when the file has every header the template requires."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return False
    try:
        worksheet = _select_sheet(workbook, template["detect"]["sheet"])
        index, _ = _read_header_index(worksheet, template["header_row"])
        return all(header_key(item) in index for item in template["detect"]["required_headers"])
    except ValueError:
        return False
    finally:
        workbook.close()


# --------------------------------------------------------------------------- reading


def read_template_dataset(
    path: Path,
    template: dict,
    *,
    borehole_code: str | None,
    template_label: str,
) -> dict:
    """Read ``path`` with ``template`` into the dataset shape the merge path consumes.

    When the template takes borehole codes from a column, only rows for
    ``borehole_code`` become intervals; every other borehole is counted and reported.
    """
    template = validate_template(template)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = _select_sheet(workbook, template["detect"]["sheet"])
        index, headers = _read_header_index(worksheet, template["header_row"])
        missing = [item for item in template["detect"]["required_headers"] if header_key(item) not in index]
        if missing:
            raise ValueError(f"File is missing required column(s): {', '.join(missing)}.")

        def position(ref: str) -> int:
            key = header_key(ref)
            if key not in index:
                raise ValueError(f'Column "{ref}" was not found in header row {template["header_row"]}.')
            return index[key]

        depth = template["depth"]
        from_col = position(depth["from"])
        to_col = position(depth["to"]) if "to" in depth else None
        thickness_col = position(depth["thickness"]) if "thickness" in depth else None
        field_cols: dict[str, int | list[int]] = {}
        for key, ref in template["fields"].items():
            field_cols[key] = [position(item) for item in ref] if isinstance(ref, list) else position(ref)

        borehole_rule = template["borehole"]["normalize"]
        borehole_col = position(template["borehole"]["column"]) if template["borehole"]["source"] == "column" else None
        rqd_divisor = 100 if template["units"]["rqd"] == "percent" else 1
        stop_after = template["end"]["stop_after_blank_rows"]

        intervals: list[dict] = []
        seams: list[dict] = []
        per_borehole: dict[str, int] = {}
        skipped_rows = 0
        dictionary_review: dict[str, int] = {}
        blank_run = 0

        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=template["data_start_row"], values_only=True),
            start=template["data_start_row"],
        ):
            if not any(clean(value) is not None for value in row):
                blank_run += 1
                if blank_run >= stop_after:
                    break
                continue
            blank_run = 0

            def cell(col: int | None):
                return row[col] if col is not None and col < len(row) else None

            if borehole_col is not None:
                file_code = normalize_borehole_code(cell(borehole_col), borehole_rule)
                if file_code is None:
                    skipped_rows += 1
                    continue
                per_borehole[file_code] = per_borehole.get(file_code, 0) + 1
                if borehole_code is not None and not borehole_codes_match(file_code, borehole_code, borehole_rule):
                    continue

            lith_from = to_float(cell(from_col))
            if to_col is not None:
                lith_to = to_float(cell(to_col))
                thickness = round(lith_to - lith_from, 3) if lith_from is not None and lith_to is not None else None
            else:
                thickness = to_float(cell(thickness_col))
                lith_to = round(lith_from + thickness, 3) if lith_from is not None and thickness is not None else None
            lithology = clean(cell(field_cols["lithology"]))
            if lith_from is None or lith_to is None or not lithology or thickness is None or thickness <= 0:
                skipped_rows += 1
                continue

            normalized = normalize_lithology(lithology)
            if normalized["needs_review"]:
                dictionary_review[normalized["code"]] = dictionary_review.get(normalized["code"], 0) + 1

            def field(key: str):
                col = field_cols.get(key)
                if col is None:
                    return None
                if isinstance(col, list):
                    parts = [str(clean(cell(item))) for item in col if clean(cell(item)) is not None]
                    return " / ".join(parts) or None
                return clean(cell(col))

            rqd_source = field("rqd")
            rqd_value = to_float(rqd_source)
            seam_name = field("seam_name")
            interval = {
                "id": f"lith-{row_number}",
                "sourceRow": row_number,
                "fromDepth": lith_from,
                "toDepth": lith_to,
                "thickness": thickness,
                "recovery": to_float(field("recovery")),
                "lithologySource": lithology,
                "lithologyCode": normalized["code"],
                "lithologyLabel": normalized["label"],
                "displayColor": normalized["color"],
                "grainSize": field("grain_size"),
                "loggedColor": field("logged_color"),
                "structuralFeatures": field("structural_features"),
                "coreDip": field("core_dip"),
                "seamName": seam_name,
                "rqd": round(rqd_value / rqd_divisor, 4) if rqd_value is not None else None,
                "rqdSource": rqd_source,
                "rqdPieceLengths": None,
                "remark": field("remark"),
            }
            intervals.append(interval)
            if seam_name:
                seams.append(
                    {
                        "id": f"seam-{row_number}",
                        "sourceRow": row_number,
                        "name": str(seam_name),
                        "fromDepth": lith_from,
                        "toDepth": lith_to,
                        "thickness": thickness,
                        "lithologyCode": normalized["code"],
                        "lithologyLabel": normalized["label"],
                    }
                )
    finally:
        workbook.close()

    min_depth = min((item["fromDepth"] for item in intervals), default=None)
    max_depth = max((item["toDepth"] for item in intervals), default=None)
    warnings: list[str] = []
    if borehole_col is not None and borehole_code is not None and not intervals:
        warnings.append(
            f'No rows matched borehole "{borehole_code}". Codes in file: '
            f'{", ".join(sorted(per_borehole)) or "none"}.'
        )
    if borehole_col is not None and borehole_code is not None:
        other_rows = sum(count for code, count in per_borehole.items() if not borehole_codes_match(code, borehole_code, borehole_rule))
        if other_rows:
            warnings.append(f"{other_rows} row(s) belong to other boreholes and were not imported here.")
    if skipped_rows:
        warnings.append(f"{skipped_rows} row(s) skipped: missing depth, lithology or borehole code, or zero thickness.")

    profile = {
        "parser": "excel_template",
        "message": f'Read with registry template "{template_label}".',
        "workbook": path.name,
        "template": {"key": template_label, "label": template_label},
        "template_key": template_label,
        "headers": headers,
        "summary": {
            "lithology_interval_count": len(intervals),
            "seam_interval_count": len(seams),
            "min_depth": min_depth,
            "max_depth": max_depth,
            "borehole_row_counts": per_borehole,
            "skipped_rows": skipped_rows,
            "dictionary_review_count": sum(dictionary_review.values()),
            "dictionary_review_codes": dictionary_review,
        },
        # Top-level depth keys feed the UI's "Depth" fact; omit them when nothing matched.
        **({"min_depth": min_depth, "max_depth": max_depth} if intervals else {}),
        "sample_rows": [
            {
                "source_row": item["sourceRow"],
                "from_depth": item["fromDepth"],
                "to_depth": item["toDepth"],
                "lithology_source": item["lithologySource"],
                "normalized_code": item["lithologyCode"],
                "seam_name": item["seamName"],
            }
            for item in intervals[:SAMPLE_ROW_LIMIT]
        ],
        "warnings": warnings,
    }
    return {
        "profile": profile,
        "borehole": {
            "totalDepth": max_depth or 0,
            "sourceSheet": template["detect"]["sheet"],
        },
        "runIntervals": [],
        "lithologyIntervals": intervals,
        "seamIntervals": seams,
    }
